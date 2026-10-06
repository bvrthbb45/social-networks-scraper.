import io
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import models as m
from app import reports
from tests import factories as f
from tests.test_auth_api import api, auth, db, enrol, env, make_user  # noqa: F401
from tests.test_import_api import session_for
from tests.test_learning import World, keys  # noqa: F401

NOW = datetime.now(timezone.utc)


def person(db, name="חייל בדיקה"):
    """A soldier whose name is encrypted exactly as the importer stores it."""
    import uuid

    from app.ingest.service import soldier_ctx
    from app.security import crypto

    sid = uuid.uuid4()
    s = m.Soldier(
        id=sid,
        personal_number_hash=uuid.uuid4().hex * 2,
        personal_number_enc=b"x",
        full_name_enc=crypto.encrypt_text(name, soldier_ctx(sid, "full_name")),
    )
    db.add(s)
    db.flush()
    return s


def decide(db, fnd, decision, hours_after, reason=None):
    db.add(
        m.Review(
            finding_id=fnd.id,
            decision=decision,
            reason=reason,
            decided_at=fnd.created_at + timedelta(hours=hours_after),
        )
    )
    fnd.status = decision


def seeded(db):
    w = World(db)
    spec = [  # kind, severity, decision, hours to decide, dismissal reason
        ("codename", "high", "confirmed", 2, None),
        ("codename", "high", "dismissed", 4, "common_word"),
        ("codename", "medium", "dismissed", 6, "common_word"),
        ("uniform", "low", "dismissed", 8, "not_relevant"),
        ("location", "high", "escalated", 10, None),
        ("codename", "medium", None, 0, None),  # undecided -> backlog
    ]
    out = []
    for i, (kind, sev, decision, hrs, reason) in enumerate(spec):
        fnd = w.finding(f"קטע {i}", kind=kind, severity=sev, key=f"k{i}")
        fnd.created_at = NOW - timedelta(days=2, hours=1)
        if decision:
            decide(db, fnd, decision, hrs, reason)
        out.append(fnd)
    old = w.finding("ישן", kind="codename", key="old")
    old.created_at = NOW - timedelta(days=200)  # outside the window
    db.commit()
    return w, out


def test_summary_counts_decisions_and_response_times(db):
    seeded(db)
    s = reports.summary(db, 30)
    assert s["findings_created"] == 6 and s["created_by_kind"] == {
        "codename": 4,
        "uniform": 1,
        "location": 1,
    }
    assert s["created_by_severity"] == {"high": 3, "medium": 2, "low": 1}
    assert s["decisions"] == {"confirmed": 1, "dismissed": 3, "escalated": 1}
    assert s["dismissal_reasons"] == {"common_word": 2, "not_relevant": 1}
    assert s["median_hours_to_decision"] == 6.0  # decisions after 2,4,6,8,10 hours
    assert (
        s["backlog"] == 2
    )  # the undecided one in the window and the old undecided one
    assert s["oldest_open_hours"] > 190 * 1  # the 200-day-old finding
    assert s["false_alarm_by_kind"]["codename"] == {
        "decided": 3,
        "dismissed": 2,
        "rate": 0.667,
    }
    assert (
        s["false_alarm_by_kind"]["uniform"]["rate"] == 1.0
        and s["false_alarm_by_kind"]["location"]["rate"] == 0.0
    )
    assert s["learning"] is None
    assert reports.summary(db, 1)["findings_created"] == 0  # window respected


def test_summary_counts_security_events_from_the_audit_trail(db):
    for action, n in (
        ("login.failed", 3),
        ("account.locked", 1),
        ("device.revoked", 2),
        ("finding.viewed", 9),
    ):
        for _ in range(n):
            db.add(m.AuditLog(action=action))
    old = m.AuditLog(action="login.failed")
    old.created_at = NOW - timedelta(days=90)
    db.add(old)
    db.commit()
    ev = reports.summary(db, 30)["security_events"]
    assert (
        ev["login.failed"] == 3
        and ev["account.locked"] == 1
        and ev["device.revoked"] == 2
    )
    assert "finding.viewed" not in ev  # only security-relevant events


def test_summary_holds_no_names_handles_or_text(db):
    seeded(db)
    blob = str(reports.summary(db, 30)) + str(reports.retention_status(db))
    for secret in ("acct_main", "קטע", "נשר"):
        assert secret not in blob


def test_consents_expiring_lists_who_needs_action(db):
    today = date.today()
    s1 = person(db)
    soon = f.consent(
        db,
        s1,
        valid_from=today - timedelta(days=300),
        valid_until=today + timedelta(days=10),
        document_ref="SOON",
    )
    lapsed = f.consent(
        db,
        person(db),
        valid_from=today - timedelta(days=300),
        valid_until=today - timedelta(days=2),
        document_ref="LAPSED",
    )
    f.consent(
        db,
        person(db),
        valid_from=today,
        valid_until=today + timedelta(days=300),
        document_ref="FAR",
    )
    revoked = f.consent(
        db,
        person(db),
        valid_from=today - timedelta(days=9),
        valid_until=today + timedelta(days=5),
        document_ref="REV",
    )
    revoked.status = "revoked"
    f.account(db, soon, status="open")
    db.commit()
    out = {c["ref"]: c for c in reports.consents_expiring(db, 30)}
    assert set(out) == {
        "SOON",
        "LAPSED",
    }  # far-future and revoked ones are not action items
    assert (
        out["SOON"]["days_left"] == 10
        and out["SOON"]["accounts"] == 1
        and out["LAPSED"]["days_left"] == -2
    )


def test_retention_status_shows_what_is_due_and_when_it_last_ran(db):
    w = World(db)
    f.post(db, w.account, delete_after=NOW - timedelta(days=1))
    f.post(db, w.account, delete_after=NOW + timedelta(days=3))
    f.post(db, w.account, delete_after=NOW + timedelta(days=60))
    db.add(m.AuditLog(action="maintenance.daily", details={"purged.posts": 4}))
    db.commit()
    r = reports.retention_status(db)
    assert (r["posts_total"], r["overdue"], r["due_within_7_days"]) == (3, 1, 1)
    assert r["last_run"]["counts"] == {"purged.posts": 4} and r["retention_days"] == 90


# --- API -------------------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "path,allowed",
    [
        ("/reports/summary", {"admin", "auditor"}),
        ("/reports/retention", {"admin", "auditor"}),
        ("/reports/consents", {"admin"}),
        ("/reports/export/findings", {"admin", "reviewer"}),
    ],
)
@pytest.mark.parametrize("role", ["admin", "auditor", "reviewer", "uploader"])
def test_report_roles(api, db, role, path, allowed):
    h = session_for(api, db, role)
    assert (api.get(path, headers=h).status_code == 200) == (role in allowed)


def test_reports_require_authentication_and_validate_the_window(api, db):
    for p in (
        "/reports/summary",
        "/reports/retention",
        "/reports/consents",
        "/reports/export/findings",
    ):
        assert api.get(p).status_code == 401
    h = session_for(api, db, "admin")
    assert api.get("/reports/summary", headers=h, params={"days": 0}).status_code == 422
    assert (
        api.get("/reports/summary", headers=h, params={"days": 366}).status_code == 422
    )


def sheet(resp, name=None):
    wb = load_workbook(io.BytesIO(resp.content))
    return wb[name] if name else wb.active, wb


def test_export_contains_only_the_requested_states_and_is_audited_without_content(
    api, db
):
    seeded(db)
    h = session_for(api, db, "reviewer")
    r = api.get(
        "/reports/export/findings", headers=h, params={"states": "confirmed,escalated"}
    )
    assert (
        r.status_code == 200
        and "spreadsheetml" in r.headers["content-type"]
        and "attachment" in r.headers["content-disposition"]
    )
    assert (
        "no-store" in r.headers["cache-control"]
        and "findings-report" in r.headers["content-disposition"]
    )
    ws, wb = sheet(r)
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3 and {row[6] for row in rows[1:]} == {
        "confirmed",
        "escalated",
    }  # header + 2
    assert all(row[2] == "acct_main" for row in rows[1:])
    info = {
        row[0]: row[1]
        for row in wb["מידע"].iter_rows(values_only=True)
        if row[1] is not None
    }
    assert (
        info["מצבים"] == "confirmed, escalated"
        and info["מספר שורות"] == 2
        and info["הופק על ידי"]
    )  # watermark
    log = db.scalars(
        select(m.AuditLog).where(m.AuditLog.action == "report.exported")
    ).one()
    assert log.details == {
        "report": "findings",
        "rows": 2,
        "states": ["confirmed", "escalated"],
        "days": 30,
    }
    assert "acct_main" not in str(log.details)


def test_export_neutralises_spreadsheet_formulas(api, db):
    w = World(db)
    fnd = w.finding('=HYPERLINK("http://evil","x")', kind="codename", key="inj")
    fnd.created_at = NOW - timedelta(hours=1)
    decide(db, fnd, "confirmed", 1)
    db.commit()
    ws, _ = sheet(
        api.get("/reports/export/findings", headers=session_for(api, db, "admin"))
    )
    snippet = list(ws.iter_rows(values_only=True))[1][8]
    assert snippet.startswith("'=") and not snippet.startswith("=")
    for dangerous in ("=1+1", "+1", "-1", "@SUM(A1)", "\tcmd", "\rcmd"):
        assert reports.safe_cell(dangerous).startswith("'")
    assert reports.safe_cell("שלום") == "שלום" and reports.safe_cell(3.5) == 3.5


@pytest.mark.parametrize(
    "states", ["", "new", "confirmed,hacked", "confirmed;dismissed"]
)
def test_export_rejects_unknown_states(api, db, states):
    h = session_for(api, db, "admin")
    assert (
        api.get(
            "/reports/export/findings", headers=h, params={"states": states}
        ).status_code
        == 422
    )


def test_consents_report_names_people_only_for_admins(api, db):
    s = person(db)
    f.consent(
        db,
        s,
        valid_from=date.today() - timedelta(days=300),
        valid_until=date.today() + timedelta(days=5),
        document_ref="X1",
    )
    db.commit()
    assert (
        api.get(
            "/reports/consents", headers=session_for(api, db, "auditor")
        ).status_code
        == 403
    )
    rows = api.get("/reports/consents", headers=session_for(api, db, "admin")).json()
    assert (
        rows[0]["ref"] == "X1"
        and rows[0]["days_left"] == 5
        and rows[0]["soldier"] == "חייל בדיקה"
    )
