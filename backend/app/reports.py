"""Aggregate reports for the information-security department, and the audited export.

Reports are numbers: no names, no post text, no handles. The two things that do identify people
(the consent-expiry list and the findings export) are restricted to roles that need them and are
audited.
"""

import io
import statistics
from collections import Counter
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import settings
from .ingest.service import soldier_ctx
from .learning import service as learning
from .models import (
    Account,
    AuditLog,
    Consent,
    Finding,
    Import,
    Post,
    Review,
    Soldier,
    User,
)
from .security import crypto

SECURITY_EVENTS = (
    "login.failed",
    "login.blocked_locked",
    "account.locked",
    "2fa.failed",
    "refresh.reuse_detected",
    "device.revoked",
    "user.credentials_reset",
    "import.rejected",
    "posts.refused",
)
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _aware(dt: datetime) -> datetime:
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _counts(db: Session, col, where=None) -> dict[str, int]:
    q = select(col, func.count()).group_by(col)
    if where is not None:
        q = q.where(where)
    return {k: n for k, n in db.execute(q).all()}


def summary(db: Session, days: int, now: datetime | None = None) -> dict:
    now = now or _now()
    since = now - timedelta(days=days)
    created = Finding.created_at >= since

    first_decision: dict = {}
    for fid, at in db.execute(
        select(Review.finding_id, func.min(Review.decided_at)).group_by(
            Review.finding_id
        )
    ).all():
        first_decision[fid] = _aware(at)
    hours = [
        (first_decision[fid] - _aware(c)).total_seconds() / 3600
        for fid, c in db.execute(
            select(Finding.id, Finding.created_at).where(created)
        ).all()
        if fid in first_decision
    ]

    latest: dict = {}
    for r in db.scalars(select(Review).order_by(Review.decided_at)).all():
        latest[r.finding_id] = r.decision
    by_kind_decided: dict[str, Counter] = {}
    for fid, kind in db.execute(select(Finding.id, Finding.kind).where(created)).all():
        if fid in latest:
            by_kind_decided.setdefault(kind, Counter())[latest[fid]] += 1
    false_alarm = {
        k: {
            "decided": sum(c.values()),
            "dismissed": c["dismissed"],
            "rate": round(c["dismissed"] / sum(c.values()), 3),
        }
        for k, c in by_kind_decided.items()
    }

    open_where = Finding.status.in_(("new", "in_review"))
    oldest = db.scalar(select(func.min(Finding.created_at)).where(open_where))
    imports = db.execute(
        select(
            func.count(),
            func.coalesce(func.sum(Import.rows_total), 0),
            func.coalesce(func.sum(Import.rows_rejected), 0),
        ).where(Import.uploaded_at >= since)
    ).one()
    events = _counts(
        db,
        AuditLog.action,
        (AuditLog.created_at >= since) & AuditLog.action.in_(SECURITY_EVENTS),
    )
    active = learning.active_model(db)
    return {
        "days": days,
        "findings_created": db.scalar(
            select(func.count()).select_from(Finding).where(created)
        )
        or 0,
        "created_by_kind": _counts(db, Finding.kind, created),
        "created_by_severity": _counts(db, Finding.severity, created),
        "decisions": _counts(db, Review.decision, Review.decided_at >= since),
        "dismissal_reasons": _counts(
            db, Review.reason, (Review.decided_at >= since) & Review.reason.is_not(None)
        ),
        "false_alarm_by_kind": false_alarm,
        "median_hours_to_decision": (
            round(statistics.median(hours), 1) if hours else None
        ),
        "backlog": db.scalar(
            select(func.count()).select_from(Finding).where(open_where)
        )
        or 0,
        "oldest_open_hours": (
            round((now - _aware(oldest)).total_seconds() / 3600, 1) if oldest else None
        ),
        "accounts_by_status": _counts(db, Account.status),
        "imports": {
            "files": imports[0],
            "rows": int(imports[1]),
            "rejected_rows": int(imports[2]),
        },
        "security_events": {k: events.get(k, 0) for k in SECURITY_EVENTS},
        "learning": (
            {
                "version": active.version,
                "age_days": (now - _aware(active.created_at)).days,
            }
            if active
            else None
        ),
    }


def consents_expiring(db: Session, days: int, today: date | None = None) -> list[dict]:
    """Active consents that end within ``days`` (or already ended and were not yet processed)."""
    today = today or date.today()
    out = []
    for c in db.scalars(
        select(Consent)
        .where(
            Consent.status == "active",
            Consent.valid_until <= today + timedelta(days=days),
        )
        .order_by(Consent.valid_until)
    ).all():
        s = db.get(Soldier, c.soldier_id)
        out.append(
            {
                "consent_id": str(c.id),
                "soldier": crypto.decrypt_text(
                    s.full_name_enc, soldier_ctx(s.id, "full_name")
                ),
                "ref": c.document_ref,
                "valid_until": c.valid_until.isoformat(),
                "days_left": (c.valid_until - today).days,
                "accounts": db.scalar(
                    select(func.count())
                    .select_from(Account)
                    .where(Account.consent_id == c.id)
                )
                or 0,
            }
        )
    return out


def retention_status(db: Session, now: datetime | None = None) -> dict:
    now = now or _now()
    last = db.scalar(
        select(AuditLog)
        .where(AuditLog.action == "maintenance.daily")
        .order_by(AuditLog.id.desc())
    )
    oldest = db.scalar(select(func.min(Post.collected_at)))
    return {
        "retention_days": settings.retention_days,
        "posts_total": db.scalar(select(func.count()).select_from(Post)) or 0,
        "overdue": db.scalar(
            select(func.count()).select_from(Post).where(Post.delete_after <= now)
        )
        or 0,
        "due_within_7_days": db.scalar(
            select(func.count())
            .select_from(Post)
            .where(
                Post.delete_after > now, Post.delete_after <= now + timedelta(days=7)
            )
        )
        or 0,
        "oldest_post_days": (now - _aware(oldest)).days if oldest else None,
        "last_run": (
            {"at": last.created_at.isoformat(), "counts": last.details}
            if last
            else None
        ),
    }


# --- export -------------------------------------------------------------------------------- #


def safe_cell(value):
    """Spreadsheet formula injection guard: text that starts like a formula is stored as plain text."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def export_findings(
    db: Session, user: User, statuses: tuple[str, ...], days: int
) -> tuple[bytes, int]:
    since = _now() - timedelta(days=days)
    latest: dict = {}
    for r in db.scalars(select(Review).order_by(Review.decided_at)).all():
        latest[r.finding_id] = r
    rows = db.execute(
        select(Finding, Post, Account)
        .join(Post, Post.id == Finding.post_id)
        .join(Account, Account.id == Post.account_id)
        .where(Finding.status.in_(statuses), Finding.created_at >= since)
        .order_by(Finding.created_at.desc())
    ).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "ממצאים"
    ws.sheet_view.rightToLeft = True
    head = [
        "זמן זיהוי",
        "רשת",
        "חשבון",
        "סוג",
        "חומרה",
        "ציון",
        "מצב",
        "סיבה",
        "קטע",
        "הוחלט על ידי",
        "מועד החלטה",
    ]
    ws.append(head)
    for f, _, a in rows:
        ev = (
            crypto.decrypt_json(f.evidence_enc, f"findings.evidence:{f.id}")
            if f.evidence_enc
            else {}
        )
        r = latest.get(f.id)
        who = db.get(User, r.reviewer_id) if r and r.reviewer_id else None
        ws.append(
            [
                safe_cell(x)
                for x in (
                    f.created_at.strftime("%Y-%m-%d %H:%M"),
                    a.platform,
                    a.username,
                    f.kind,
                    f.severity,
                    float(f.score),
                    f.status,
                    crypto.decrypt_text(f.reason_enc, f"findings.reason:{f.id}"),
                    ev.get("snippet") or "",
                    (who.display_name or who.email) if who else "",
                    r.decided_at.strftime("%Y-%m-%d %H:%M") if r else "",
                )
            ]
        )
    for i, w in enumerate((17, 11, 20, 14, 9, 7, 11, 48, 40, 18, 17), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    info = wb.create_sheet("מידע")
    info.sheet_view.rightToLeft = True
    for line in (
        ["חומר רגיש – לשימוש מחלקת ביטחון מידע בלבד"],
        ["הופק על ידי", safe_cell(user.display_name or user.email)],
        ["מועד הפקה", _now().strftime("%Y-%m-%d %H:%M UTC")],
        ["מצבים", ", ".join(statuses)],
        ["טווח (ימים)", days],
        ["מספר שורות", len(rows)],
        ["ההפקה נרשמה ביומן הביקורת."],
    ):
        info.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)
