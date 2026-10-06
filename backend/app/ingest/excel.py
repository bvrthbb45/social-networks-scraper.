"""Parse and validate the roster workbook. Pure functions: no database, no network.

Safety rules
* only real .xlsx (zip + OOXML), no macros, bounded size / rows / decompressed size;
* formulas are never evaluated (``data_only`` + read-only mode);
* columns this system has no use for (notably PHONE NUMBERS) are never read, only counted by
  name, so a file that contains them does not leak them into the database or logs;
* rejection reasons are codes with row numbers: never the row content.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from urllib.parse import urlparse

from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ROWS = 20_000

PLATFORMS = ("instagram", "tiktok", "facebook")

# Canonical field -> accepted header spellings (compared after normalisation).
HEADERS: dict[str, tuple[str, ...]] = {
    "personal_number": (
        "מספר אישי",
        "מ.א",
        'מ"א',
        "personal number",
        "personal_number",
    ),
    "full_name": ("שם מלא", "שם", "full name", "full_name", "name"),
    "unit": ("יחידה", "unit"),
    "platform": ("פלטפורמה", "רשת", "רשת חברתית", "platform", "network"),
    "account": (
        "חשבון",
        "שם משתמש",
        "משתמש",
        "קישור",
        "username",
        "handle",
        "url",
        "account",
    ),
    "status": ("סטטוס", "מצב", "פתוח/סגור", "status"),
    "consent_ref": (
        "אסמכתת הסכמה",
        "מסמך הסכמה",
        "אסמכתא",
        "consent ref",
        "consent_ref",
    ),
    "consent_signed": ("תאריך חתימה", "נחתם ב", "consent signed", "signed_on"),
    "consent_from": ("תוקף מ", "תוקף מתאריך", "valid from", "valid_from"),
    "consent_until": ("תוקף עד", "תוקף עד תאריך", "valid until", "valid_until"),
}
REQUIRED = (
    "personal_number",
    "full_name",
    "platform",
    "account",
    "status",
    "consent_ref",
)

STATUS_WORDS = {
    "open": ("פתוח", "פתוחה", "ציבורי", "open", "public"),
    "closed": ("סגור", "סגורה", "פרטי", "closed", "private"),
    "unknown": ("", "לא ידוע", "unknown"),
    "not_found": ("לא נמצא", "לא קיים", "not found", "not_found"),
}
PLATFORM_WORDS = {
    "instagram": ("instagram", "אינסטגרם", "אינסטה", "ig"),
    "tiktok": ("tiktok", "טיקטוק", "טיק טוק"),
    "facebook": ("facebook", "פייסבוק", "fb"),
}
HOSTS = {
    "instagram.com": "instagram",
    "tiktok.com": "tiktok",
    "facebook.com": "facebook",
    "fb.com": "facebook",
}
_PHONE_HEADER = re.compile(
    r"טלפון|נייד|פלאפון|phone|mobile|cell|whatsapp|וואטסאפ|ווטסאפ", re.I
)
_USERNAME_OK = re.compile(r"^[a-z0-9._]{1,100}$")


class FileRejected(Exception):
    """The whole file is unacceptable. ``code`` is safe to show."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass
class Row:
    line: int
    personal_number: str
    full_name: str
    unit: str | None
    platform: str
    username: str
    status: str
    consent_ref: str
    consent_signed: date
    consent_from: date
    consent_until: date


@dataclass
class Parsed:
    rows: list[Row] = field(default_factory=list)
    rejected: dict[str, list[int]] = field(
        default_factory=dict
    )  # reason code -> line numbers
    total: int = 0
    ignored_columns: int = (
        0  # how many columns were not read (count only, never names/content)
    )
    phone_columns_ignored: int = 0

    def reject(self, code: str, line: int) -> None:
        self.rejected.setdefault(code, []).append(line)


def _norm_header(v) -> str:
    s = re.sub(r"[\s_\-]+", " ", str(v or "")).strip().casefold()
    return s.replace("'", "").replace("״", '"')


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # Excel stores 1234567 as 1234567.0
        v = int(v)
    return str(v).strip()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_status(v) -> str | None:
    s = _text(v).casefold()
    for canon, words in STATUS_WORDS.items():
        if s in words:
            return canon
    return None


def parse_platform(v, account_value: str = "") -> str | None:
    s = _text(v).casefold()
    for canon, words in PLATFORM_WORDS.items():
        if s in words:
            return canon
    if s == "" and account_value.startswith(
        "http"
    ):  # platform can be inferred from a URL
        return _platform_from_url(account_value)
    return None


def _platform_from_url(value: str) -> str | None:
    host = (urlparse(value).hostname or "").lower()
    host = host[4:] if host.startswith("www.") else host
    host = host[2:] if host.startswith("m.") else host
    return HOSTS.get(host)


def parse_username(value: str) -> str | None:
    """Accept '@name', 'name' or a profile URL; return the normalised username or None."""
    s = _text(value)
    if s.startswith("http"):
        parts = [p for p in urlparse(s).path.split("/") if p]
        if not parts:
            return None
        s = parts[0]
        if s in ("profile.php", "people", "pages", "p", "reel", "stories", "explore"):
            return None  # not a stable username URL
    s = s.lstrip("@").strip().lower()
    return s if _USERNAME_OK.match(s) else None


def _check_container(data: bytes) -> None:
    if len(data) > MAX_BYTES:
        raise FileRejected("file_too_large")
    if not data.startswith(b"PK"):
        raise FileRejected("not_xlsx")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            if "[Content_Types].xml" not in names or not any(
                n.startswith("xl/") for n in names
            ):
                raise FileRejected("not_xlsx")
            if any("vbaProject" in n for n in names):
                raise FileRejected("macros_not_allowed")
            if any(n.startswith("xl/externalLinks/") for n in names):
                raise FileRejected("external_links_not_allowed")
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise FileRejected("file_too_large")  # decompression bomb
    except zipfile.BadZipFile:
        raise FileRejected("not_xlsx")


def _open(data: bytes):
    _check_container(data)
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise FileRejected("unreadable")


def _map_headers(header_row) -> tuple[dict[str, int], int, int]:
    index: dict[str, int] = {}
    ignored = phones = 0
    lookup = {_norm_header(a): f for f, aliases in HEADERS.items() for a in aliases}
    for i, h in enumerate(header_row):
        n = _norm_header(h)
        if n in lookup and lookup[n] not in index:
            index[lookup[n]] = i
        elif n:
            ignored += 1
            if _PHONE_HEADER.search(n):
                phones += 1
    return index, ignored, phones


def parse_roster(data: bytes, today: date | None = None) -> Parsed:
    today = today or date.today()
    wb = _open(data)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise FileRejected("empty")
    index, ignored, phones = _map_headers(header)
    missing = [f for f in REQUIRED if f not in index]
    if missing:
        raise FileRejected("missing_columns:" + ",".join(missing))
    out = Parsed(ignored_columns=ignored, phone_columns_ignored=phones)

    def cell(r, f):
        i = index.get(f)
        return r[i] if i is not None and i < len(r) else None

    for line, r in enumerate(rows_iter, start=2):
        if not any(c not in (None, "") for c in r):
            continue  # blank line
        out.total += 1
        if out.total > MAX_ROWS:
            raise FileRejected("too_many_rows")
        number, name = _text(cell(r, "personal_number")), _text(cell(r, "full_name"))
        account_raw = _text(cell(r, "account"))
        if not number or not name or not re.fullmatch(r"[0-9A-Za-z\-]{5,15}", number):
            out.reject("bad_identity", line)
            continue
        platform = parse_platform(cell(r, "platform"), account_raw)
        if platform is None:
            # "whatsapp" and everything else that is not a supported public platform
            out.reject("platform_not_allowed", line)
            continue
        username = parse_username(account_raw)
        if username is None:
            out.reject("bad_account", line)
            continue
        status = parse_status(cell(r, "status"))
        if status is None:
            out.reject("bad_status", line)
            continue
        ref = _text(cell(r, "consent_ref"))
        if not ref:
            out.reject(
                "no_consent", line
            )  # no consent reference = the row is not imported
            continue
        signed, vfrom, vuntil = (
            _date(cell(r, "consent_signed")),
            _date(cell(r, "consent_from")),
            _date(cell(r, "consent_until")),
        )
        if signed is None or vfrom is None or vuntil is None:
            out.reject("bad_consent_dates", line)
            continue
        if vuntil < vfrom or signed > today:
            out.reject("bad_consent_dates", line)
            continue
        if not (vfrom <= today <= vuntil):
            out.reject("consent_not_in_force", line)
            continue
        out.rows.append(
            Row(
                line,
                number,
                name,
                _text(cell(r, "unit")) or None,
                platform,
                username,
                status,
                ref[:120],
                signed,
                vfrom,
                vuntil,
            )
        )
    wb.close()
    return out


# --- watch-list workbook ------------------------------------------------------------------ #

TERM_HEADERS = {
    "term": ("מונח", "שם קוד", "שם", "term", "codename"),
    "aliases": ("כינויים", "שמות נוספים", "aliases"),
    "kind": ("סוג", "kind"),
    "severity": ("חומרה", "severity"),
}
KIND_WORDS = {
    "codename": ("שם קוד", "codename"),
    "site": ("אתר", "מיקום", "site", "location"),
    "unit": ("יחידה", "unit"),
    "other": ("", "אחר", "other"),
}
SEVERITY_WORDS = {
    "low": ("נמוכה", "low"),
    "medium": ("בינונית", "medium", ""),
    "high": ("גבוהה", "high"),
}


@dataclass
class Term:
    line: int
    term: str
    aliases: list[str]
    kind: str
    severity: str


def _word(v, table) -> str | None:
    s = _text(v).casefold()
    for canon, words in table.items():
        if s in words:
            return canon
    return None


def parse_terms(data: bytes) -> tuple[list[Term], dict[str, list[int]]]:
    wb = _open(data)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise FileRejected("empty")
    lookup = {_norm_header(a): f for f, al in TERM_HEADERS.items() for a in al}
    index: dict[str, int] = {}
    for i, h in enumerate(header):
        f = lookup.get(_norm_header(h))
        if f and f not in index:
            index[f] = i
    if "term" not in index:
        raise FileRejected("missing_columns:term")
    terms: list[Term] = []
    rejected: dict[str, list[int]] = {}
    total = 0
    for line, r in enumerate(rows_iter, start=2):
        if not any(c not in (None, "") for c in r):
            continue
        total += 1
        if total > MAX_ROWS:
            raise FileRejected("too_many_rows")

        def g(f):
            i = index.get(f)
            return r[i] if i is not None and i < len(r) else None

        term = _text(g("term"))
        kind, sev = _word(g("kind"), KIND_WORDS), _word(g("severity"), SEVERITY_WORDS)
        if len(term) < 2 or len(term) > 120:
            rejected.setdefault("bad_term", []).append(line)
        elif kind is None or sev is None:
            rejected.setdefault("bad_kind_or_severity", []).append(line)
        else:
            aliases = [
                a.strip() for a in re.split(r"[,;\n]", _text(g("aliases"))) if a.strip()
            ]
            terms.append(Term(line, term, aliases[:20], kind, sev))
    wb.close()
    return terms, rejected
